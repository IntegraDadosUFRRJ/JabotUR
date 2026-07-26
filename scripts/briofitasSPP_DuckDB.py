"""
Como a limpeza foi feita:
  - Parcela "4*" -> 4 (asterisco tratado como marcação de erro/observação)
  - "Parcela", "Amostra" e "Substrato" seguem o padrão de preenchimento apenas na 1ª linha do grupo: recebem fill-down 
  - "Substrato" é multivalorado (ex.: linha 132 "A + S "): tabela externa substrato + associativa briofita_substrato
  - "Observações" também é multivalorada, separada por ";": tabela observacao + associativa briofita_observacao
"""

import re
import openpyxl
import duckdb

XLSX_PATH = "./briófitas JB_UFRRJ - Oliveira_Santos.xlsx"
DB_PATH = "./briofitas_teste.duckdb"

# === LEGENDAS ===
COORDENADAS_PARCELA = {
    "1": "22°45'54.2\"S 43°41'31.1\"O",   
    "2": "22°45'57.3\"S 43°41'34\"O",     
    "3": "22°45'54.3\"S 43°41'33.7\"O",   
    "4": "22°45'54.9\"S 43°41'34.5\"O",
    "aleat": "Amostragem aleatória (sem coordenada fixa)",
}

SUBSTRATO_NOMES = {
    "RZ": ("Raiz de árvore", "RZ"),
    "S": ("Solo", "S"),
    "TD": ("Tronco em decomposição", "TD"),
    "TV": ("Tronco vivo", "TV"),
    "A": ("Artificial", "A"),
}

IDENTIFICACAO_LEGENDA = {
    1: "Identificada",
    2: "Parcialmente Identificada",
    3: "Dúvida",
}

FILO_NOMES = {
    "M": "Musgos",
    "H": "Hepáticas",
}


wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
ws = wb["SPP - JB"]

raw_rows = []
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 12)]
    if all(v is None for v in vals):
        continue
    raw_rows.append(vals)

records = []
last_parcela = None
last_amostra = None
last_substrato = None

for (parcela, amostra, status, filo, familia, genero, especie,
     autor, forma, substrato, obs) in raw_rows:

    # fill-down
    if parcela is not None:
        last_parcela = parcela
    if amostra is not None:
        last_amostra = amostra
    if substrato is not None:
        last_substrato = substrato

    parcela_val = last_parcela
    amostra_val = last_amostra
    forma_val = forma
    substrato_val = last_substrato

    parcela_str = str(parcela_val).strip()
    if parcela_str == "4*":
        parcela_str = "4"
    elif parcela_str.endswith(".0"):
        parcela_str = parcela_str[:-2]

    records.append({
        "parcela": parcela_str,
        "amostra": int(amostra_val) if amostra_val is not None else None,
        "status": int(status),
        "filo": filo,
        "familia": familia,
        "genero": genero,
        "especie": especie,
        "autor": autor,
        "forma": forma_val,
        "substrato_raw": substrato_val,
        "obs_raw": obs,
    })

print(f"Linhas lidas da planilha (com dados): {len(raw_rows)}")
print(f"Registros processados: {len(records)}")

# === TABELAS ===
def norm(s):
    return s.strip() if isinstance(s, str) else s

# filo 
filos = sorted({r["filo"] for r in records})
filo_id = {sigla: i + 1 for i, sigla in enumerate(filos)}
filo_rows = [(filo_id[s], FILO_NOMES.get(s, s), s) for s in filos]

# familia (filha do filo pai)
familia_to_filo = {}
for r in records:
    familia_to_filo[r["familia"]] = r["filo"]
familias = sorted(familia_to_filo)
familia_id = {f: i + 1 for i, f in enumerate(familias)}
familia_rows = [(familia_id[f], filo_id[familia_to_filo[f]], f) for f in familias]

# genero (filha da familia pai)
genero_to_familia = {}
for r in records:
    genero_to_familia[r["genero"]] = r["familia"]
generos = sorted(genero_to_familia)
genero_id = {g: i + 1 for i, g in enumerate(generos)}
genero_rows = [(genero_id[g], familia_id[genero_to_familia[g]], g) for g in generos]

# especie (filha do genero pai)
especie_to_genero = {}
especie_to_autor = {}
for r in records:
    especie_to_genero[r["especie"]] = r["genero"]
    especie_to_autor.setdefault(r["especie"], r["autor"])
especies = sorted(especie_to_genero)
especie_id = {e: i + 1 for i, e in enumerate(especies)}

# autor
autores = sorted({r["autor"] for r in records})
autor_id = {a: i + 1 for i, a in enumerate(autores)}
autor_rows = [(autor_id[a], a) for a in autores]
especie_rows = [(especie_id[e], genero_id[especie_to_genero[e]], autor_id[especie_to_autor[e]], e) for e in especies]

# identificacao 
identificacao_rows = [(k, v) for k, v in IDENTIFICACAO_LEGENDA.items()]

# forma_de_vida 
formas = sorted({r["forma"] for r in records if r["forma"] is not None})
forma_id = {f: i + 1 for i, f in enumerate(formas)}
forma_rows = [(forma_id[f], f) for f in formas]

# parcela
parcelas = sorted({r["parcela"] for r in records})
parcela_id = {p: i + 1 for i, p in enumerate(parcelas)}
parcela_rows = [(parcela_id[p], COORDENADAS_PARCELA.get(p, "não informado")) for p in parcelas]

# substrato - multivalores separados por: "+"
substrato_codes = set()
for r in records:
    if r["substrato_raw"]:
        parts = [norm(p) for p in r["substrato_raw"].split("+")]
        substrato_codes.update(parts)
substrato_codes = sorted(substrato_codes)
substrato_id = {c: i + 1 for i, c in enumerate(substrato_codes)}
substrato_rows = [
    (substrato_id[c], SUBSTRATO_NOMES.get(c, (c, c))[0], SUBSTRATO_NOMES.get(c, (c, c))[1])
    for c in substrato_codes
]

# observacao - multivalores separados por: ";"
obs_texts = set()
for r in records:
    if r["obs_raw"]:
        parts = [norm(p) for p in r["obs_raw"].split(";")]
        obs_texts.update(parts)
obs_texts = sorted(obs_texts)
observacao_id = {o: i + 1 for i, o in enumerate(obs_texts)}
observacao_rows = [(observacao_id[o], o, o[:10]) for o in obs_texts]


coleta_rows = []
briofita_substrato_rows = []
briofita_observacao_rows = []

for i, r in enumerate(records, start=1):
    coleta_rows.append((
        i,
        especie_id[r["especie"]],
        forma_id[r["forma"]] if r["forma"] is not None else None,
        r["status"],
        parcela_id[r["parcela"]],
        r["amostra"],
    ))
    if r["substrato_raw"]:
        for p in r["substrato_raw"].split("+"):
            briofita_substrato_rows.append((i, substrato_id[norm(p)]))
    if r["obs_raw"]:
        for p in r["obs_raw"].split(";"):
            briofita_observacao_rows.append((i, observacao_id[norm(p)]))

# === BANCO DUCKDB === 
con = duckdb.connect(DB_PATH)
con.execute("PRAGMA enable_progress_bar=false")

con.execute("""
    CREATE OR REPLACE TABLE filo (
        id INTEGER PRIMARY KEY,
        nome VARCHAR,
        sigla VARCHAR(1)
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE familia (
        id INTEGER PRIMARY KEY,
        id_filo INTEGER REFERENCES filo(id),
        nome VARCHAR
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE genero (
        id INTEGER PRIMARY KEY,
        id_familia INTEGER REFERENCES familia(id),
        nome VARCHAR
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE autor (
        id INTEGER PRIMARY KEY,
        nome VARCHAR
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE especie (
        id INTEGER PRIMARY KEY,
        id_genero INTEGER REFERENCES genero(id),
        id_autor INTEGER REFERENCES autor(id),
        nome VARCHAR
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE identificacao (
        id INTEGER PRIMARY KEY,
        descricao VARCHAR
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE forma_vida (
        id INTEGER PRIMARY KEY,
        tipo VARCHAR
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE parcela (
        id INTEGER PRIMARY KEY,
        coordenada VARCHAR
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE substrato (
        id INTEGER PRIMARY KEY,
        nome VARCHAR,
        sigla VARCHAR(3)
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE observacao (
        id INTEGER PRIMARY KEY,
        descricao VARCHAR,
        abreviacao VARCHAR
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE coleta (
        id INTEGER PRIMARY KEY,
        id_especie INTEGER REFERENCES especie(id),
        id_forma INTEGER REFERENCES forma_vida(id),
        id_identificacao INTEGER REFERENCES identificacao(id),
        id_parcela INTEGER REFERENCES parcela(id),
        amostra INTEGER
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE briofita_substrato (
        id_briofita INTEGER REFERENCES coleta(id),
        id_substrato INTEGER REFERENCES substrato(id),
        PRIMARY KEY (id_briofita, id_substrato)
    )
""")
con.execute("""
    CREATE OR REPLACE TABLE briofita_observacao (
        id_briofita INTEGER REFERENCES coleta(id),
        id_observacao INTEGER REFERENCES observacao(id),
        PRIMARY KEY (id_briofita, id_observacao)
    )
""")

con.executemany("INSERT INTO filo VALUES (?, ?, ?)", filo_rows)
con.executemany("INSERT INTO familia VALUES (?, ?, ?)", familia_rows)
con.executemany("INSERT INTO genero VALUES (?, ?, ?)", genero_rows)
con.executemany("INSERT INTO autor VALUES (?, ?)", autor_rows)
con.executemany("INSERT INTO especie VALUES (?, ?, ?, ?)", especie_rows)
con.executemany("INSERT INTO identificacao VALUES (?, ?)", identificacao_rows)
con.executemany("INSERT INTO forma_vida VALUES (?, ?)", forma_rows)
con.executemany("INSERT INTO parcela VALUES (?, ?)", parcela_rows)
con.executemany("INSERT INTO substrato VALUES (?, ?, ?)", substrato_rows)
con.executemany("INSERT INTO observacao VALUES (?, ?, ?)", observacao_rows)
con.executemany("INSERT INTO coleta VALUES (?, ?, ?, ?, ?, ?)", coleta_rows)
con.executemany("INSERT INTO briofita_substrato VALUES (?, ?)", briofita_substrato_rows)
con.executemany("INSERT INTO briofita_observacao VALUES (?, ?)", briofita_observacao_rows)

# === VERIFICAÇÃO ===
print("\n--- Contagem de linhas por tabela ---")
for t in ["filo", "familia", "genero", "especie", "autor", "identificacao",
          "forma_vida", "parcela", "substrato", "observacao", "coleta",
          "briofita_substrato", "briofita_observacao"]:
    n = con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
    print(f"{t:22s} {n}")

print("\nAmostra (5 linhas)")
sample = con.execute("""
    SELECT c.id, fi.sigla AS filo, fa.nome AS familia, g.nome AS genero,
        e.nome AS especie, au.nome AS autor, fv.tipo AS forma_vida,
        p.coordenada AS parcela, c.amostra, ident.descricao AS status
    FROM coleta c
    JOIN especie e ON e.id = c.id_especie
    JOIN genero g ON g.id = e.id_genero
    JOIN familia fa ON fa.id = g.id_familia
    JOIN filo fi ON fi.id = fa.id_filo
    JOIN autor au ON au.id = e.id_autor
    LEFT JOIN forma_vida fv ON fv.id = c.id_forma
    JOIN parcela p ON p.id = c.id_parcela
    JOIN identificacao ident ON ident.id = c.id_identificacao
    ORDER BY c.id
    LIMIT 5
""").fetchdf()
print(sample.to_string(index=False))

print("\n Substratos multivalorados")
multi = con.execute("""
    SELECT id_briofita, COUNT(*) AS n_substratos
    FROM briofita_substrato
    GROUP BY id_briofita
    HAVING COUNT(*) > 1
    ORDER BY id_briofita
""").fetchdf()
print(multi.to_string(index=False))

con.close()
print(f"\nBanco de teste salvo em: {DB_PATH}")